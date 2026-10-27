#!/usr/bin/env python3
"""Client presentation, exact shaped type outlines, vector PDF and SVG PowerPoint.

Run from the repo root. Requirements: fonttools, uharfbuzz, Pillow, cairosvg,
python-pptx and pypdf. No external images, fonts or network assets in the output.
"""
import copy
import datetime
import hashlib
import html
import io
import json
import math
from pathlib import Path
import sys
import zipfile

import cairosvg
import uharfbuzz as hb
from fontTools.ttLib import TTFont
from fontTools.pens.svgPathPen import SVGPathPen
from fontTools.pens.boundsPen import BoundsPen
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Link
from pptx import Presentation
from pptx.util import Inches
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement

OUT = Path('docs/presentation')
BUILD = Path('build/presentation')
W, H = 1600, 900
INK, PAPER, PERI = '#191A2C', '#F5F1E9', '#6868AC'
PALE, MUTED, CORAL = '#DFDCEE', '#626275', '#EFAC8B'
DATE = '2026-10-27'
DOWNLOAD = 'https://github.com/rahatulyerbolat05-dotcom/Very-Peri-serif/raw/refs/heads/claude/upbeat-faraday-mzabid/dist/VeryPeriSerif-latest.zip'
PROJECT = 'https://github.com/rahatulyerbolat05-dotcom/Very-Peri-serif'
FONT_PATHS = {'R':Path('fonts/ttf/VeryPeriSerif-Regular.ttf'),
              'I':Path('fonts/ttf/VeryPeriSerif-Italic.ttf'),
              'S':Path('/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf')}


class Face:
    def __init__(self, path, key):
        self.key, self.path = key, path
        self.font = TTFont(path); self.gs = self.font.getGlyphSet()
        self.upm = self.font['head'].unitsPerEm
        self.hf = hb.Font(hb.Face(path.read_bytes()))
        self.hf.scale = (self.upm, self.upm)
        self.paths, self.boxes = {}, {}

    def shape(self, text, features=None, lang='ru'):
        b = hb.Buffer(); b.add_str(text); b.guess_segment_properties(); b.language=lang
        hb.shape(self.hf,b,features or {})
        assert all(g.codepoint for g in b.glyph_infos), ('Missing glyph',self.key,text)
        result=[]
        for g,p in zip(b.glyph_infos,b.glyph_positions):
            name=self.font.getGlyphName(g.codepoint)
            if name not in self.paths:
                pen=SVGPathPen(self.gs); self.gs[name].draw(pen)
                self.paths[name]=pen.getCommands()
                box=BoundsPen(self.gs); self.gs[name].draw(box)
                self.boxes[name]=box.bounds
            result.append((name,p.x_advance,p.x_offset,p.y_offset))
        return result


FACES={key:Face(path,key) for key,path in FONT_PATHS.items()}


class Canvas:
    def __init__(self,title,bg=PAPER,w=W,h=H,number=None):
        self.title,self.bg,self.w,self.h=title,bg,w,h
        self.uid='p'+hashlib.sha256(title.encode()).hexdigest()[:10]
        self.parts=[f'<rect width="{w}" height="{h}" fill="{bg}"/>']
        self.used=set(); self.records=[]
        self.number=number

    def rect(self,x,y,w,h,fill,rx=0,stroke=None,sw=1,opacity=1):
        self.parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{rx}" fill="{fill}" stroke="{stroke or "none"}" stroke-width="{sw}" opacity="{opacity}"/>')

    def line(self,x1,y1,x2,y2,color,sw=1,opacity=1):
        self.parts.append(f'<path d="M{x1},{y1} L{x2},{y2}" stroke="{color}" stroke-width="{sw}" opacity="{opacity}"/>')

    def circle(self,cx,cy,r,fill,stroke=None,sw=1):
        self.parts.append(f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="{fill}" stroke="{stroke or "none"}" stroke-width="{sw}"/>')

    def text(self,text,x,y,size=100,face='R',color=INK,align='left',tracking=0,
             features=None,lang='ru',maxw=None,crop=False,opacity=1):
        font=FACES[face]; shaped=font.shape(text,features,lang)
        width=sum(p[1] for p in shaped)/font.upm*size+max(0,len(shaped)-1)*tracking
        if maxw is not None and width>maxw:
            size*=maxw/width
            width=maxw
            assert size>=12,('Text became too small',text,size)
        scale=size/font.upm
        if align=='center':x-=width/2
        elif align=='right':x-=width
        start=x; boxes=[]; chunks=[]
        for i,(name,advance,dx,dy) in enumerate(shaped):
            self.used.add((face,name))
            if font.paths[name]:
                px,py=x+dx*scale,y-dy*scale
                chunks.append(f'<use xlink:href="#{self.uid}_{face}_{name}" transform="translate({px:.5f} {py:.5f}) scale({scale:.8f} {-scale:.8f})"/>')
                b=font.boxes[name]
                boxes.append((px+b[0]*scale,py-b[3]*scale,px+b[2]*scale,py-b[1]*scale))
            x+=advance*scale+(tracking if i<len(shaped)-1 else 0)
        if boxes:
            box=[min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)]
            if not crop:
                assert box[0]>=-0.1 and box[1]>=-0.1 and box[2]<=self.w+.1 and box[3]<=self.h+.1,('Text outside slide',self.title,text,box)
        else:box=[start,y,start+width,y]
        self.parts.append(f'<g fill="{color}" opacity="{opacity}" aria-label="{html.escape(text,quote=True)}"><title>{html.escape(text)}</title>'+''.join(chunks)+'</g>')
        self.records.append({'text':text,'face':face,'size':round(size,2),'bounds':[round(v,3) for v in box],'features':features or {},'intentional_crop':crop})
        return width

    def para(self,lines,x,y,size=23,color=INK,leading=34,face='S',maxw=None):
        for i,line in enumerate(lines):self.text(line,x,y+i*leading,size,face,color,maxw=maxw)

    def hero_stack(self,lines,x,top,bottom,size,maxw,gap=20,lang='ru'):
        """Space large mixed-face lines by their real ink, including descenders."""
        metrics=[]
        for text,face,color in lines:
            font=FACES[face];glyphs=font.shape(text,lang=lang)
            width=sum(p[1] for p in glyphs)/font.upm*size
            z=min(size,size*maxw/width)
            boxes=[(font.boxes[n][1]+dy,font.boxes[n][3]+dy) for n,a,dx,dy in glyphs if font.boxes[n]]
            lo=min(p[0] for p in boxes);hi=max(p[1] for p in boxes)
            metrics.append((text,face,color,z,lo,hi,font.upm))
        ink=sum((hi-lo)*z/upm for t,f,c,z,lo,hi,upm in metrics)
        factor=min(1,(bottom-top-gap*(len(lines)-1))/ink)
        y=top
        for text,face,color,z,lo,hi,upm in metrics:
            z*=factor
            self.text(text,x,y+hi*z/upm,z,face,color,lang=lang)
            y+=(hi-lo)*z/upm+gap

    def furniture(self,section,light=False,right_light=False):
        color=PAPER if light else INK
        right_color=PAPER if right_light else color
        self.text('VERY PERI SERIF',84,63,16,'S',color,tracking=1.6)
        self.text(section,1516,63,16,'S',right_color,align='right',tracking=1.0)
        self.text('YERBOLAT RAKHATULY / DEW',84,852,13,'S',color,tracking=.8)
        if self.number is not None:self.text(f'{self.number:02d} / 14',1516,852,13,'S',right_color,align='right',tracking=.8)

    def svg(self):
        defs=[]
        for face,name in sorted(self.used):
            defs.append(f'<path id="{self.uid}_{face}_{name}" d="{FACES[face].paths[name]}"/>')
        return f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{self.w}" height="{self.h}" viewBox="0 0 {self.w} {self.h}" role="img"><title>{html.escape(self.title)}</title><defs>'+''.join(defs)+'</defs>'+''.join(self.parts)+'</svg>'


def slides():
    pages=[]
    c=Canvas('Very Peri Serif — презентация гарнитуры',PERI,number=1)
    c.furniture('ПРЕЗЕНТАЦИЯ ГАРНИТУРЫ',True)
    c.text('Very Peri',80,397,410,'R',PAPER,maxw=1420)
    c.text('Serif',1460,732,430,'I',PAPER,align='right',maxw=760)
    c.line(88,563,481,563,PAPER,1,opacity=.45)
    c.para(['Акцидентная антиква.','Русский, қазақша, Latin.'],88,620,24,PAPER,38)
    c.text('REGULAR / ITALIC',88,755,17,'S',PAPER,tracking=1.2)
    pages.append(c)

    c=Canvas('Характер в каждой букве',PAPER,number=2);c.furniture('ХАРАКТЕР')
    c.hero_stack([('Характер','R',INK),('в каждой','I',PERI),('букве.','R',INK)],84,110,790,248,960,20)
    c.line(1080,154,1080,764,INK,1,.2)
    c.para(['Узкие пропорции.','Выразительный контраст.','Тонкие связи и живые','окончания штрихов.'],1130,235,23,leading=36,maxw=365)
    c.para(['Для айдентики, обложек,','афиш и крупных','редакционных заголовков.'],1130,510,23,leading=36,maxw=365)
    c.text('01',1130,727,90,'R',PERI)
    c.text('гарнитура / два начертания',1240,722,17,'S',MUTED,maxw=270)
    pages.append(c)

    c=Canvas('Два начертания — два голоса',PAPER,number=3)
    c.rect(800,0,800,900,INK);c.furniture('ДВА НАЧЕРТАНИЯ',right_light=True)
    for x,face,col,label in ((84,'R',INK,'REGULAR'),(880,'I',PAPER,'ITALIC')):
        c.text(label,x,185,19,'S',col,tracking=2)
        c.text('Внимание',x,395,205,face,col,maxw=630)
        c.text('к деталям.',x,590,200,face,col,maxw=630)
    c.para(['Собранный ритм.','Спокойная опора композиции.'],84,730,22,INK,32)
    c.para(['Рукописное движение.','Акцент внутри той же системы.'],880,730,22,PAPER,32,maxw=620)
    pages.append(c)

    c=Canvas('Әсемдік әріптен басталады',PERI,number=4);c.furniture('ҚАЗАҚША',True)
    c.hero_stack([('Әсемдік','R',PAPER),('әріптен','I',PAPER),('басталады.','R',PAPER)],84,110,790,270,1060,20,lang='kk')
    c.line(1190,150,1190,755,PAPER,1,.3)
    for i,text in enumerate(('Әә  Ғғ  Ққ','Ңң  Өө  Ұұ','Үү  Һһ  Іі')):
        c.text(text,1240,265+i*155,78,'R',PAPER,maxw=285,lang='kk')
    c.para(['Казахские буквы —','часть гарнитуры','в обоих начертаниях.'],1240,690,19,PAPER,28,maxw=285)
    pages.append(c)

    c=Canvas('Детали кириллицы',PALE,number=5);c.furniture('КИРИЛЛИЦА / ДЕТАЛИ')
    c.text('Узнать по одной букве.',84,196,135,maxw=1430)
    labels=['Регулярная д','Курсивная ж','Регулярная я','Курсивная ң']
    for i,(ch,face) in enumerate((('д','R'),('ж','I'),('я','R'),('ң','I'))):
        cx=260+360*i
        if i:c.line(cx-180,275,cx-180,740,INK,1,.15)
        c.text(ch,cx,595,650,face,INK,align='center',maxw=285,lang='kk')
        c.text(labels[i],cx,795,18,'S',INK,align='center',maxw=315)
    pages.append(c)

    c=Canvas('Альтернативная л — своё написание',PAPER,number=6);c.furniture('ВАРИАНТЫ НАПИСАНИЯ')
    c.text('Своё написание.',84,207,160,maxw=1370)
    c.rect(84,265,695,500,PALE);c.rect(801,265,715,500,PERI)
    c.text('л',430,653,650,'I',INK,align='center')
    c.text('л',1160,653,650,'I',PAPER,align='center',features={'ss10':1})
    c.text('ОСНОВНАЯ ФОРМА',116,310,15,'S',INK,tracking=1.2)
    c.text('ОКРУГЛАЯ АРКА И ХВОСТ',833,310,15,'S',PAPER,tracking=1.2)
    c.text('лилия · люблю',420,730,95,'I',INK,align='center',maxw=580)
    c.text('лилия · люблю',1158,730,95,'I',PAPER,align='center',maxw=580,features={'ss10':1})
    c.text('Включаемая альтернатива в Italic / стилистический набор ss10',84,803,19,'S',MUTED,maxw=1420)
    pages.append(c)

    c=Canvas('Цифры и знак тенге',INK,number=7);c.furniture('ЦИФРЫ И ЗНАКИ',True)
    c.text('12 500 ₸',84,375,350,'R',PAPER,maxw=1430)
    c.line(84,467,1516,467,PAPER,1,.25)
    c.text('0123456789',84,625,170,'I',PALE,maxw=1030)
    c.para(['Табличные цифры','для ровных колонок.'],1160,565,20,PAPER,31,maxw=350)
    c.text('111 000',1158,685,86,'R',PAPER,features={'tnum':1},maxw=350)
    c.text('888 000',1158,760,86,'R',PAPER,features={'tnum':1},maxw=350)
    c.text('₸  ₽  €  $',84,773,105,'R',PAPER,maxw=730)
    pages.append(c)

    c=Canvas('Пластика строки и лигатуры',CORAL,number=8);c.furniture('ПЛАСТИКА СТРОКИ')
    c.text('Плавность',84,275,225,maxw=1000)
    c.text('в деталях.',84,464,240,'I',INK,maxw=1000)
    c.text('office',1510,693,315,'R',INK,align='right',maxw=990,features={'liga':1})
    c.line(84,553,440,553,INK,1,.35)
    c.para(['Стандартные лигатуры','поддерживают движение','соседних букв.'],84,620,22,INK,33,maxw=400)
    c.text('fi',1280,348,245,'I',INK,align='center',features={'liga':1})
    c.text('affine · final · film',1510,770,76,'I',INK,align='right',maxw=1020,features={'liga':1})
    pages.append(c)

    c=Canvas('Обложка — пример применения',PAPER,number=9);c.furniture('ПРИМЕНЕНИЕ / ОБЛОЖКА')
    c.text('Первое',84,320,205,maxw=700)
    c.text('впечатление.',84,510,205,'I',PERI,maxw=700)
    c.para(['Обложка эссе.','Название становится','главным изображением.'],84,665,23,INK,35,maxw=630)
    c.rect(931,145,505,660,'#D4CEC6')
    c.rect(908,124,505,660,PERI)
    c.rect(908,124,13,660,'#4E4E87')
    c.text('ЛИТЕРАТУРНАЯ СЕРИЯ / 01',953,182,13,'S',PAPER,tracking=.7,maxw=410)
    c.text('Культура',950,375,153,'R',PAPER,maxw=425)
    c.text('паузы',950,525,185,'I',PAPER,maxw=425)
    c.line(953,620,1367,620,PAPER,1,.45)
    c.para(['О внимании к простому','и пространстве для мысли.'],953,670,20,PAPER,31,maxw=405)
    c.text('КОНЦЕПТ ОБЛОЖКИ',953,744,12,'S',PAPER,tracking=1,maxw=410)
    pages.append(c)

    c=Canvas('Афиша — пример применения',PERI,number=10);c.furniture('ПРИМЕНЕНИЕ / АФИША',True)
    c.text('Голос',78,382,365,'R',PAPER,maxw=1320)
    c.text('формы',1500,713,385,'I',PAPER,align='right',maxw=1330)
    c.line(86,482,603,482,PAPER,1,.45)
    c.para(['ТИПОГРАФИЧЕСКАЯ','ВЫСТАВКА / КОНЦЕПТ'],86,550,18,PAPER,28,maxw=500)
    c.text('2026',1510,176,86,'R',PAPER,align='right')
    pages.append(c)

    c=Canvas('Редакционная веб-страница — пример применения',PALE,number=11);c.furniture('ПРИМЕНЕНИЕ / DIGITAL')
    c.rect(84,123,1432,640,PAPER)
    c.rect(84,123,1432,44,'#D0CDDF')
    for x in (107,128,149):c.circle(x,145,5,'#A09CAF')
    c.text('VERY PERI JOURNAL',122,217,15,'S',INK,tracking=1)
    c.text('КУЛЬТУРА      ДИЗАЙН      СЛОВА',1475,217,13,'S',INK,align='right',tracking=.5)
    c.line(122,245,1475,245,INK,1,.2)
    c.text('Слово',124,410,188,'R',INK,maxw=780)
    c.text('обретает форму',124,566,184,'I',PERI,maxw=800)
    c.para(['Внимание к букве меняет ощущение от всей страницы.','Заголовок задаёт тон, а паузы помогают его услышать.'],127,646,20,INK,32,maxw=825)
    c.text('ЧИТАТЬ МАТЕРИАЛ',127,724,13,'S',INK,tracking=.8)
    c.parts.append(f'<path d="M302,719 H325 M319,713 L325,719 L319,725" fill="none" stroke="{INK}" stroke-width="1.2"/>')
    c.rect(1035,284,397,438,PERI)
    c.text('а',1226,648,685,'I',PAPER,align='center',maxw=350)
    c.text('НАЗВАНИЯ И МАКЕТЫ — КОНЦЕПТЫ ПРИМЕНЕНИЯ',84,803,14,'S',MUTED,tracking=.8)
    pages.append(c)

    c=Canvas('Слова, которые остаются',INK,number=12);c.furniture('VERY PERI SERIF',True)
    c.text('Слова,',84,310,280,'R',PAPER,maxw=1415)
    c.text('которые',84,520,280,'I',PALE,maxw=1415)
    c.text('остаются.',84,752,280,'R',PAPER,maxw=1415)
    pages.append(c)

    c=Canvas('Кириллический набор',PAPER,number=13);c.furniture('ПРИЛОЖЕНИЕ / НАБОР')
    c.text('Русский. Қазақша.',84,174,112,maxw=1415)
    chars='АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯӘҒҚҢӨҰҮҺІ'
    assert len(chars)==42
    for i,ch in enumerate(chars):
        col,row=i%7,i//7
        color=PERI if i>=33 else INK
        c.text(ch+ch.lower(),183+205*col,280+94*row,90,'R',color,align='center',maxw=170,lang='kk')
    c.line(84,791,1516,791,INK,1,.18)
    c.text('Также: исходная латиница Instrument Serif, цифры и знаки.',84,825,18,'S',MUTED,maxw=1415)
    # Furniture footer would clash with this last note; the note sits above it.
    pages.append(c)

    c=Canvas('Авторство, лицензия и файлы',PAPER,number=14);c.furniture('ПРИЛОЖЕНИЕ / О ПРОЕКТЕ')
    c.text('Very Peri',84,294,235,'R',INK,maxw=730)
    c.text('Serif',84,507,265,'I',PERI,maxw=730)
    c.text('REGULAR / ITALIC',88,602,18,'S',INK,tracking=1.3)
    c.text('Версия 1.114',88,645,21,'S',MUTED)
    c.rect(84,705,545,71,INK)
    c.text('СКАЧАТЬ ГАРНИТУРУ',111,750,18,'S',PAPER,tracking=1,maxw=420)
    c.parts.append(f'<path d="M567,741 H601 M593,733 L601,741 L593,749" fill="none" stroke="{PAPER}" stroke-width="1.5"/>')
    c.para(['Кириллическое расширение Instrument Serif.','Автор кириллицы и руководитель разработки:','Yerbolat Rakhatuly (DEW).'],900,210,22,INK,34,maxw=620)
    c.para(['Латинская основа — Instrument Serif.','Rodrigo Fuenzalida / Jordan Egstad.','Строчная д Regular — на основе Cormorant','Garamond, The Cormorant Project Authors.'],900,370,19,MUTED,31,maxw=620)
    c.para(['SIL Open Font License 1.1.','TTF / OTF — для установки. WOFF2 — для сайта.','Кириллица создана с ИИ-инструментами','под руководством автора.'],900,561,19,INK,31,maxw=620)
    c.para(['Макеты применения — презентационные концепты.','Служебные подписи — Noto Sans.'],900,737,15,MUTED,25,maxw=620)
    pages.append(c)
    assert len(pages)==14
    return pages


def posters():
    result=[]
    c=Canvas('Very Peri Serif / Голос формы',PERI,1080,1350)
    c.text('VERY PERI SERIF',70,90,19,'S',PAPER,tracking=2)
    c.text('Голос',64,482,355,'R',PAPER,maxw=945)
    c.text('формы',1010,823,355,'I',PAPER,align='right',maxw=945)
    c.line(74,1031,1006,1031,PAPER,1,.4)
    c.para(['Русский / Қазақша / Latin','Regular / Italic'],74,1110,26,PAPER,43)
    c.text('YERBOLAT RAKHATULY / DEW',74,1273,16,'S',PAPER,tracking=1)
    result.append(('01-golos-formy',c))
    c=Canvas('Very Peri Serif / Қазақша',PAPER,1080,1350)
    c.text('VERY PERI SERIF',70,90,19,'S',INK,tracking=2)
    c.text('Әсемдік',65,430,287,'R',PERI,maxw=950,lang='kk')
    c.text('әріптен',65,682,287,'I',INK,maxw=950,lang='kk')
    c.text('басталады.',65,942,273,'R',PERI,maxw=950,lang='kk')
    c.line(74,1075,1006,1075,INK,1,.2)
    c.text('Ә Ғ Қ Ң Ө Ұ Ү Һ І',74,1174,85,'R',INK,maxw=935,lang='kk')
    c.text('YERBOLAT RAKHATULY / DEW',74,1273,16,'S',INK,tracking=1)
    result.append(('02-asemdik',c))
    c=Canvas('Very Peri Serif / Альтернативная л',INK,1080,1350)
    c.text('VERY PERI SERIF',70,90,19,'S',PAPER,tracking=2)
    c.text('л',538,1032,1450,'I',PALE,align='center',features={'ss10':1},maxw=900)
    c.text('ДЕТАЛЬ С ХАРАКТЕРОМ',74,1195,24,'S',PAPER,tracking=1.1,maxw=935)
    c.text('ITALIC / SS10',74,1254,16,'S',PAPER,tracking=1)
    c.text('DEW',1006,1254,16,'S',PAPER,align='right',tracking=1)
    result.append(('03-el-alternate',c))
    return result


def make_pdf(pages):
    writer=PdfWriter()
    for i,c in enumerate(pages,1):
        pdf=cairosvg.svg2pdf(bytestring=c.svg().encode())
        (BUILD/f'{i:02d}.pdf').write_bytes(pdf)
        writer.add_page(PdfReader(io.BytesIO(pdf)).pages[0])
    # SVG artwork remains vector; the final download link is a real PDF link.
    writer.add_annotation(len(pages)-1,Link(rect=(84*.75,(900-776)*.75,629*.75,(900-705)*.75),url=DOWNLOAD))
    writer.add_metadata({'/Title':'Very Peri Serif — презентация гарнитуры',
        '/Author':'Yerbolat Rakhatuly (DEW)',
        '/Subject':'Regular / Italic · Русский / Қазақша / Latin · 1.114',
        '/CreationDate':"D:20261027000000+05'00'"})
    with (OUT/'VeryPeriSerif-Presentation.pdf').open('wb') as f:writer.write(f)


def make_pptx(pages):
    prs=Presentation(); prs.slide_width=Inches(40/3);prs.slide_height=Inches(7.5)
    prs.core_properties.title='Very Peri Serif — презентация гарнитуры'
    prs.core_properties.author='Yerbolat Rakhatuly (DEW)'
    prs.core_properties.subject='Regular / Italic · Русский / Қазақша / Latin'
    moment=datetime.datetime(2026,10,27,tzinfo=datetime.timezone.utc)
    prs.core_properties.created=moment;prs.core_properties.modified=moment
    for i,c in enumerate(pages,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_picture(str(OUT/'previews'/f'{i:02d}.png'),0,0,width=prs.slide_width,height=prs.slide_height)
        svg=Part(PackURI(f'/ppt/media/vector{i:02d}.svg'),'image/svg+xml',slide.part.package,c.svg().encode())
        rid=slide.part.relate_to(svg,RT.IMAGE)
        blip=shape._element.blipFill.blip
        extlst=OxmlElement('a:extLst');ext=OxmlElement('a:ext')
        ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        from lxml import etree
        child=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',
                            nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
        child.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',rid)
        ext.append(child);extlst.append(ext);blip.append(extlst)
        slide.notes_slide.notes_text_frame.text=(c.title+'\n\nVery Peri Serif 1.114. Кириллица: Yerbolat Rakhatuly (DEW).\n'
            'Слайд содержит векторный SVG и PNG для совместимости. Редактируемый вектор — в папке slides.\n'
            'Латинская основа: Instrument Serif. Макеты применения — концепты.')
    prs.save(OUT/'VeryPeriSerif-Presentation.pptx')


def contact_sheet(pages):
    tw,th=400,225;gap=20;rows=math.ceil(len(pages)/4)
    image=Image.new('RGB',(4*tw+5*gap,rows*(th+40)+gap), '#D8D6D2')
    draw=ImageDraw.Draw(image);font=ImageFont.truetype(str(FONT_PATHS['S']),15)
    for i,c in enumerate(pages):
        im=Image.open(OUT/'previews'/f'{i+1:02d}.png').convert('RGB');im=im.resize((tw,th),Image.Resampling.LANCZOS)
        x=gap+(i%4)*(tw+gap);y=gap+(i//4)*(th+40)
        image.paste(im,(x,y));draw.text((x,y+th+6),f'{i+1:02d}  {c.title[:42]}',font=font,fill=INK)
    image.save(OUT/'overview.png')


def viewer(pages):
    title='Very Peri Serif — презентация'
    frames=''.join(f'<section aria-label="{html.escape(c.title,quote=True)}"'+(' class="active"' if i==0 else '')+'>'+c.svg()+'</section>' for i,c in enumerate(pages))
    code='''const pages=[...document.querySelectorAll('section')];let pos=0;function go(step){pages[pos].classList.remove('active');pos=(pos+step+pages.length)%pages.length;pages[pos].classList.add('active');document.getElementById('counter').textContent=(pos+1)+' / '+pages.length;}document.getElementById('prev').onclick=()=>go(-1);document.getElementById('next').onclick=()=>go(1);document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();go(1);}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();go(-1);}});document.getElementById('full').onclick=()=>document.documentElement.requestFullscreen?.();'''
    css='''*{box-sizing:border-box}body{margin:0;background:#11121c;color:#f5f1e9;font:14px system-ui}main{height:calc(100dvh - 64px);display:grid;place-items:center;padding:12px}section{display:none;max-width:1600px;width:min(100%,calc((100dvh - 88px)*16/9));box-shadow:0 20px 60px #0005}section.active{display:block}svg{display:block;width:100%;height:auto}nav{height:64px;display:flex;align-items:center;justify-content:center;gap:24px}button{background:transparent;border:1px solid #8b87ac;color:inherit;padding:9px 18px;border-radius:0;cursor:pointer}button:hover{background:#343149}#counter{min-width:50px;text-align:center}@media print{body{background:white}main{display:block;height:auto;padding:0}section{display:block;width:100%;break-after:page;box-shadow:none}nav{display:none}}'''
    (OUT/'presentation.html').write_text(f'<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{title}</title><style>{css}</style><main>{frames}</main><nav aria-label="Навигация презентации"><button id="prev" aria-label="Предыдущий слайд">←</button><span id="counter">1 / {len(pages)}</span><button id="next" aria-label="Следующий слайд">→</button><button id="full">На весь экран</button></nav><script>{code}</script></html>')


def package(pages,posters_):
    readme='''Very Peri Serif — презентация для заказчика
==========================================
Версия шрифта: 1.114. Автор кириллицы: Yerbolat Rakhatuly (DEW).

VeryPeriSerif-Presentation.pdf — 14 векторных слайдов, формат 16:9.
VeryPeriSerif-Presentation.pptx — те же слайды для показа в PowerPoint.
  SVG даёт векторное изображение в совместимых версиях Office; PNG служит
  резервным отображением. Текст слайдов переведён в контуры.
presentation.html — автономный просмотр в браузере. Стрелки переключают слайды.
slides/ — исходные SVG: можно открыть и менять в Figma, Illustrator, Inkscape.
previews/ — изображения слайдов 2400×1350 px.
social/ — три постера 4:5, PNG 2160×2700 px и исходные SVG.
overview.png — обзор всех слайдов.

Надписи набраны реальными файлами Very Peri Serif 1.114 и переведены в вектор
после OpenType-шейпинга. Русский/казахский, лигатуры, ss10 и tnum используются
из собранного шрифта. Установка шрифтов для просмотра не требуется.
Служебные подписи — Noto Sans. Макеты обложки, афиши и журнала — концепты.

Кириллица: Yerbolat Rakhatuly (DEW). Латинская основа — Instrument Serif,
Rodrigo Fuenzalida / Jordan Egstad. Строчная д Regular — модификация
Cormorant Garamond, The Cormorant Project Authors. См. AUTHORS.txt и OFL.txt.
Гарнитура: SIL OFL 1.1. Оригинальные материалы презентации подготовлены
для Yerbolat Rakhatuly (DEW); полные файлы шрифтов находятся в отдельном архиве.

Сборка: python tools/presentation.py (из корня репозитория).
Зависимости: fonttools, uharfbuzz, Pillow, cairosvg, python-pptx, pypdf.
Нативные приложения Office/Figma отдельно не тестировались; основная
версия для отправки — проверенный векторный PDF.
'''
    (OUT/'README.txt').write_text(readme)
    manifest={'date':DATE,'purpose':'Client-facing presentation','font_version':'1.114',
              'author':'Yerbolat Rakhatuly (DEW)','slides':len(pages),'social_posters':len(posters_),
              'slide_size':[W,H],'preview_size':[2400,1350],'social_size':[2160,2700],
              'font_sha256':{k:hashlib.sha256(p.read_bytes()).hexdigest() for k,p in FONT_PATHS.items()},
              'pages':[{'number':i+1,'title':c.title,'text_objects':c.records} for i,c in enumerate(pages)],
              'qa':{'missing_glyphs':0,'unintended_out_of_bounds':0,'all_text_outlined':True}}
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    # Repository navigation in README.md is not meaningful in the portable kit;
    # its self-contained README.txt ships instead.
    targets=[p for p in OUT.rglob('*') if p.is_file() and p.name!='README.md']
    members={p.relative_to(OUT).as_posix():p.read_bytes() for p in targets}
    members['AUTHORS.txt']=Path('AUTHORS.txt').read_bytes();members['OFL.txt']=Path('OFL.txt').read_bytes()
    members['tools/presentation.py']=Path(__file__).read_bytes()
    members['SHA256SUMS.txt']=''.join(hashlib.sha256(raw).hexdigest()+'  '+name+'\n' for name,raw in sorted(members.items())).encode()
    archive=Path('dist/VeryPeriSerif-Presentation-1.114.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for name,raw in sorted(members.items()):
            info=zipfile.ZipInfo('VeryPeriSerif-Presentation/'+name,(2026,10,27,0,0,0))
            info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o100644<<16;z.writestr(info,raw)
    with zipfile.ZipFile(archive) as z:assert z.testzip() is None
    print(archive)


def main():
    for path in (OUT,BUILD,OUT/'slides',OUT/'previews',OUT/'social'):path.mkdir(exist_ok=True,parents=True)
    for key in ('R','I'):
        assert FACES[key].font['name'].getDebugName(5)=='Version 1.114'
    ss10=FACES['I'].shape('л',{'ss10':1})
    assert ss10[0][0]=='uni043B.ss10'
    assert len(FACES['R'].shape('office',{'liga':1}))<len(FACES['R'].shape('office',{'liga':0}))
    assert len({sum(p[1] for p in FACES['R'].shape(ch,{'tnum':1})) for ch in '0123456789'})==1
    pages=slides(); posters_=posters()
    for i,c in enumerate(pages,1):
        raw=c.svg().encode();(OUT/'slides'/f'{i:02d}.svg').write_bytes(raw)
        cairosvg.svg2png(bytestring=raw,write_to=str(OUT/'previews'/f'{i:02d}.png'),output_width=2400,output_height=1350)
    for name,c in posters_:
        raw=c.svg().encode();(OUT/'social'/(name+'.svg')).write_bytes(raw)
        cairosvg.svg2png(bytestring=raw,write_to=str(OUT/'social'/(name+'.png')),output_width=2160,output_height=2700)
    make_pdf(pages);make_pptx(pages);contact_sheet(pages);viewer(pages);package(pages,posters_)
    print(json.dumps({'slides':len(pages),'posters':len(posters_),'pdf_pages':len(PdfReader(OUT/'VeryPeriSerif-Presentation.pdf').pages),'source_fonts_unchanged':True}))


if __name__=='__main__':main()
